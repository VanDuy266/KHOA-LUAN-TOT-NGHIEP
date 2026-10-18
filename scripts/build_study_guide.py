"""Tạo sổ tay học và bảo vệ từ đúng các bảng thực nghiệm hiện có."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parents[1]


def markdown_table(frame):
    def cell(value):
        return f"{value:.3f}" if isinstance(value, float) else str(value)
    lines = ["| " + " | ".join(frame.columns) + " |", "| " + " | ".join(["---"] * len(frame.columns)) + " |"]
    lines += ["| " + " | ".join(cell(value) for value in row) + " |" for row in frame.itertuples(index=False, name=None)]
    return "\n".join(lines)


def main():
    metrics = pd.read_csv(ROOT / "reports/metrics_summary.csv")
    by_pin = pd.read_csv(ROOT / "reports/metrics_by_battery.csv")
    ablation = pd.read_csv(ROOT / "reports/ablation_summary.csv")
    temporal = pd.read_csv(ROOT / "reports/temporal_metrics.csv")
    text = (ROOT / "docs/so-tay-hoc-va-bao-ve-template.md").read_text(encoding="utf-8")
    tables = {
        "{{METRICS}}": metrics,
        "{{PIN_METRICS}}": by_pin[by_pin.model.isin(["Cycle Ridge", "Random Forest"])][["held_out_battery", "model", "mae", "rmse"]],
        "{{ABLATION}}": ablation,
        "{{TEMPORAL}}": temporal[["battery", "model", "n_train", "n_test", "mae"]],
    }
    for key, frame in tables.items():
        text = text.replace(key, markdown_table(frame))
    if "{{" in text:
        raise ValueError("Sổ tay còn trường chưa điền")
    (ROOT / "docs/so-tay-hoc-va-bao-ve.md").write_text(text, encoding="utf-8")
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2)
    in_code, rows = False, []

    def flush_table():
        if not rows:
            return
        values = [[value.strip() for value in row.strip("|").split("|")] for row in rows]
        values = [row for row in values if not all(value == "---" for value in row)]
        table = doc.add_table(rows=0, cols=len(values[0]))
        table.style = "Table Grid"
        for row in values:
            for item, value in zip(table.add_row().cells, row):
                item.text = value
                for paragraph in item.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
        rows.clear()

    for line in text.splitlines():
        if line.startswith("```"):
            flush_table()
            in_code = not in_code
            continue
        if in_code:
            paragraph = doc.add_paragraph(style="No Spacing")
            run = paragraph.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(9)
        elif line.startswith("|"):
            rows.append(line)
        else:
            flush_table()
            if line.startswith("#"):
                prefix, title = line.split(" ", 1)
                doc.add_heading(title, min(len(prefix), 3))
            elif line.strip():
                doc.add_paragraph(line.replace("**", "").replace("`", ""))
    flush_table()
    doc.save(ROOT / "docs/so-tay-hoc-va-bao-ve.docx")
    print("Đã tạo sổ tay Markdown và Word từ kết quả thực nghiệm đang có.")


if __name__ == "__main__":
    main()
