"""Tạo báo cáo tiến độ và slide giữa kỳ ngày 10/11/2026 từ kết quả thực nghiệm."""
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt
from pptx.dml.color import RGBColor
from build_documents import md_table

ROOT = Path(__file__).resolve().parents[1]


def main():
    sdu = pd.read_csv(ROOT / "reports/metrics_summary.csv")
    xjtu = pd.read_csv(ROOT / "reports/xjtu/metrics_summary.csv")
    info = json.loads((ROOT / "reports/xjtu/dataset_summary.json").read_text())
    run = json.loads((ROOT / "reports/xjtu/run_metadata.json").read_text())
    rf = sdu[sdu.model == "Random Forest"].iloc[0]
    within = xjtu[xjtu.design == "XJTU held-out"]
    xrf = within[within.model == "Random Forest"].iloc[0]
    xmulti = within[within.model == "RF multi-source"].iloc[0]
    text = f'''# BÁO CÁO GIỮA KỲ KHÓA LUẬN TỐT NGHIỆP

Ngày báo cáo: 10/11/2026. Bản tiến độ cập nhật từ kết quả thực nghiệm đã lưu, không phải xác nhận nghiệm thu cuối kỳ.

Đề tài: Phân tích và dự đoán tình trạng sức khỏe pin xe điện bằng các mô hình học máy.

Sinh viên: Nguyễn Văn Duy — MSSV 2311553949 — Chuyên ngành Khoa Học Dữ Liệu.

## 1. Mục tiêu và phạm vi

Phân tích dữ liệu điện áp, dòng điện, nhiệt độ và dung lượng; dự đoán SOH từ đầu chu kỳ xả; so sánh mô hình trên pin chưa được học; xây dựng ứng dụng Vincent D. Dữ liệu thực nghiệm ở cấp tế bào trong phòng thí nghiệm. Chưa kiểm chứng trên BMS xe thực tế và chưa triển khai RUL; hai nội dung này không được ghi là chức năng đã hoàn thành.

## 2. Dữ liệu đã triển khai

SDU 2025: 86 pin Lishen 2,4 Ah, hai giai đoạn thuộc cùng các pin; 10.505.065 điểm đo được giữ, 20.823 chu kỳ xả tham chiếu. Nguồn Zenodo 14859405, CC BY 4.0.

XJTU 2024: chọn 23/55 pin Lishen NCM 2 Ah thuộc Batch-1/2 vì quy trình cố định. Đã tải 23 MAT nguyên gốc, kiểm CRC/SHA, trích xuất từ {info['raw_points']:,} điểm đo. Có {info['valid_cycles']:,} chu kỳ đạt kiểm tra đầu vào; {run['model_cycles']:,} chu kỳ sau khi không dùng 23 chu kỳ đo đầu 0,2C cho đánh giá xả 1C. Nguồn Zenodo 10963339, CC BY 4.0.

Giữ nguyên nguồn. Không ghép hai bộ như một tập đồng nhất và không sửa số dòng để làm đẹp quy mô. Dung lượng danh định và dòng xả khác nhau, nên dùng quy trình riêng và báo cáo chuyển miền riêng.

## 3. Quy trình xử lý và đánh giá

Nguồn → kiểm tra cấu trúc/dung lượng → tách chu kỳ xả → cửa sổ 120 giây → đặc trưng hoặc chuỗi → giữ riêng mã pin → huấn luyện → dự đoán pin kiểm thử → MAE/RMSE/R² → ứng dụng.

SDU dùng 12 đặc trưng. XJTU dùng 11 đặc trưng chung: bỏ phase_code, chia dòng điện cho dung lượng danh định. Dung lượng đo cả chu kỳ chỉ tạo nhãn SOH, không làm đầu vào. GRU dùng 25 vị trí V/I/T và biến ngữ cảnh; chuẩn hóa chỉ học từ train.

Chia 5 vòng theo pin vật lý, seed 42; không chia ngẫu nhiên các điểm cảm biến của cùng pin vào hai tập. MAE/RMSE tính theo điểm phần trăm SOH, kết quả chính lấy trung bình theo pin. R² có thể âm khi mô hình chuyển miền kém.

## 4. Kết quả SDU

{md_table(sdu)}

Random Forest là cấu hình demo SDU cố định: MAE {rf.mae:.3f}, RMSE {rf.rmse:.3f}, R² {rf.r2:.3f}. Điểm này thuộc kiểm thử pin giữ riêng; artifact ứng dụng sau đó refit trên toàn bộ dữ liệu, không dùng điểm refit thay điểm kiểm thử.

## 5. Kết quả bổ sung XJTU

{md_table(within.drop(columns='design'))}

RF đa nguồn học từ SDU và các pin XJTU thuộc train của từng vòng, cân bằng tổng trọng số hai nguồn; không học nhãn pin XJTU đang kiểm tra. So sánh cần xem cả MAE và RMSE; thêm dữ liệu không bảo đảm sai số giảm.

Kết quả chuyển miền và giữ riêng batch được trình bày riêng:

{md_table(xjtu[xjtu.design != 'XJTU held-out'])}

Không suy rộng kết quả này thành dự đoán chính xác mọi loại pin xe điện. Đây là thực nghiệm ban đầu trên cấu hình cố định; chưa có đánh giá đa seed hoặc kiểm chứng xe.

## 6. Sản phẩm hiện có

Vincent D có tra cứu pin, biểu đồ, so sánh mô hình, dự đoán một/lô SDU, lịch sử, kiểm tra chất lượng, phát hiện xu hướng SOH bất thường và xuất CSV. Mục XJTU tách riêng kết quả kiểm thử, nguồn, dự đoán MAT bằng RF riêng hoặc đa nguồn. Các cờ bất thường chỉ yêu cầu xem xét, chưa là chẩn đoán hỏng pin.

Mã nguồn, bảng kết quả, notebook, báo cáo và slide được lưu trên GitHub. Dữ liệu gốc lớn được tải từ nguồn bằng script, không cần đưa toàn bộ vào RAM để chạy giao diện. Các notebook đã có đầu ra để xem riêng mà không mở Streamlit.

## 7. Nội dung đã làm và phần cần tiếp tục

Đã triển khai: dữ liệu chính SDU, kiểm thử mô hình theo pin, mô hình Ridge–RF phần dư, giám sát xu hướng, ứng dụng; mở rộng XJTU Batch-1/2, kiểm thử độc lập theo pin, kiểm tra chuyển miền và RF đa nguồn.

Cần tiếp tục trước nghiệm thu cuối kỳ: rà soát diễn giải sai số theo pin/điều kiện, kiểm tra tính ổn định cấu hình, đối chiếu báo cáo và mã nguồn; chỉnh định dạng theo mẫu trường và nhận xét giảng viên. BMS trên xe và RUL là giới hạn/hướng mở rộng, không tự động đưa vào phạm vi bắt buộc.

## 8. Kế hoạch chuẩn bị ngày 10/11/2026

| Thời gian | Công việc |
| --- | --- |
| 25–29/10 | Rà soát nguồn, nhãn và phân hoạch; kiểm tra mã và chạy sản phẩm trên máy cá nhân |
| 30/10–03/11 | Phân tích sai số, đối chiếu bảng/biểu đồ và hoàn thiện báo cáo tiến độ |
| 04–07/11 | Chốt slide, lựa chọn ví dụ demo và bản dự phòng kết quả |
| 08–09/11 | Tập trình bày, chuẩn bị trả lời câu hỏi về SOH, chia dữ liệu và giới hạn |
| 10/11 | Báo cáo giữa kỳ và ghi nhận góp ý |

Lịch trên là kế hoạch chuẩn bị đề xuất; điều chỉnh theo lịch và mẫu giảng viên cung cấp.

## 9. Tài liệu trình chiếu và bàn giao giữa kỳ

Slide: docs/slide-giua-ky-10-11-2026.pptx. Báo cáo tiến độ: docs/bao-cao-giua-ky-10-11-2026.docx. Notebook: 01_eda, 02_ket_qua_tung_phan và 03_thuc_nghiem_xjtu. Thực nghiệm lưu ở reports và reports/xjtu. Demo chạy bằng chay_windows.ps1 từ thư mục gốc. Hồ sơ cuối kỳ vẫn theo yêu cầu chính thức của trường.
'''
    stem = ROOT / "docs/bao-cao-giua-ky-10-11-2026"
    stem.with_suffix(".md").write_text(text)
    doc = Document()
    doc.core_properties.author = "Nguyễn Văn Duy"
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(13)
    style.paragraph_format.line_spacing = 1.3
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            group = []
            while i < len(lines) and lines[i].startswith("|"):
                group.append([x.strip() for x in lines[i].strip("|").split("|")]); i += 1
            table = doc.add_table(rows=1, cols=len(group[0])); table.style = "Table Grid"
            for c, v in zip(table.rows[0].cells, group[0]):c.text = v
            for row in group[2:]:
                for c,v in zip(table.add_row().cells,row):c.text=v
            continue
        if line.startswith("# "):doc.add_heading(line[2:],0)
        elif line.startswith("## "):doc.add_heading(line[3:],1)
        elif line.strip():doc.add_paragraph(line)
        i += 1
    doc.save(stem.with_suffix(".docx"))
    slides = [
        ("Báo cáo giữa kỳ · 10/11/2026", ["Phân tích và dự đoán tình trạng sức khỏe pin xe điện bằng các mô hình học máy", "Nguyễn Văn Duy · MSSV 2311553949", "Chuyên ngành Khoa Học Dữ Liệu"], "Tôi trình bày tiến độ, kết quả thực nghiệm ban đầu và phần còn tiếp tục."),
        ("Bài toán và mục tiêu", ["SOH: dung lượng hiện tại so với dung lượng tham chiếu", "Dự đoán từ 120 giây đầu chu kỳ xả", "So sánh mô hình trên pin chưa được học"], "SOH khác SOC. Bài toán chính là hồi quy SOH; tôi chưa triển khai RUL."),
        ("Phạm vi nghiên cứu", ["Dữ liệu tế bào pin lithium-ion trong phòng lab", "SDU là nguồn chính, XJTU là nguồn bổ sung", "Chưa kiểm chứng bằng BMS trên xe thực tế"], "Tôi không khẳng định mô hình dùng được ngay cho mọi bộ pin xe điện."),
        ("Dữ liệu SDU", ["86 pin Lishen NMC 2,4 Ah; hai giai đoạn", "10.505.065 điểm đo · 20.823 chu kỳ", "Nguồn Zenodo 2025 · CC BY 4.0"], "Hai giai đoạn thuộc cùng 86 pin; số dòng đo không phải số mẫu chu kỳ độc lập."),
        ("Dữ liệu XJTU tương đồng", ["Lishen 18650 NCM 2 Ah · nguồn 2024", "Chọn 23 pin Batch-1/2 có quy trình cố định", f"{run['model_cycles']:,} chu kỳ xả 1C trong đánh giá chính"], "Tôi giữ nguyên MAT và kiểm tra nguồn. Chu kỳ đầu xả 0,2C được giữ để đối chiếu nhưng không dùng trong đánh giá 1C."),
        ("Xử lý và đặc trưng", ["Kiểm tra → tách chu kỳ → lấy cửa sổ đầu xả", "V/I/T tạo đầu vào; dung lượng cả chu kỳ tạo nhãn", "XJTU đổi dòng điện sang C-rate theo 2 Ah"], "Tôi không đưa dung lượng hoặc SOH thực vào đặc trưng. Dữ liệu gốc không bị sửa."),
        ("Các mô hình thực nghiệm", ["Ridge · Random Forest · LightGBM · GRU", "Mốc trung bình làm đối chứng", "Ridge–RF phần dư; RF đa nguồn là phương án thử thêm"], "Đây là thuật toán và cách kết hợp có sẵn, không tuyên bố phát minh thuật toán mới."),
        ("Thiết kế kiểm thử", ["5 vòng chia theo mã pin; train/test không trùng pin", "Chuẩn hóa chỉ học trên tập train", "MAE, RMSE, R² lấy trung bình theo pin"], "Nếu chia ngẫu nhiên các dòng cùng pin, kết quả có thể lạc quan do rò rỉ dữ liệu."),
        ("Kết quả ban đầu SDU", [f"RF: MAE {rf.mae:.3f} · RMSE {rf.rmse:.3f}", f"R² {rf.r2:.3f} · kiểm thử giữ riêng pin", "Bảng đầy đủ và sai số theo pin ở notebook/ứng dụng"], "Điểm này lấy từ dự đoán giữ riêng pin, không phải dự đoán lại tập đã học."),
        ("Kết quả bổ sung XJTU", [f"RF: MAE {xrf.mae:.3f} · RMSE {xrf.rmse:.3f}", f"RF đa nguồn: MAE {xmulti.mae:.3f} · RMSE {xmulti.rmse:.3f}", "MAE/RMSE theo điểm phần trăm SOH", "RF SDU cũ → XJTU: MAE 12,836; chưa chuyển trực tiếp"], "Thêm dữ liệu có thể tăng hoặc giảm sai số; tôi báo cáo đúng kết quả và giới hạn điều kiện đo."),
        ("Sản phẩm Vincent D", ["Tra cứu, biểu đồ, kiểm tra và dự đoán SOH", "Trang XJTU: kết quả, nguồn và dự đoán MAT riêng", "CSV kết quả; notebook và mô hình xem ngoài ứng dụng"], "Tôi demo một pin ở trang kiểm thử giữ riêng trước, sau đó giới thiệu dự đoán ứng dụng; hai kết quả có ý nghĩa khác nhau."),
        ("Kế hoạch tiếp theo", ["Rà soát sai số và độ ổn định thực nghiệm", "Hoàn thiện báo cáo theo góp ý giảng viên", "Chuẩn bị demo, bản dự phòng và nội dung bảo vệ"], "BMS trên xe và RUL vẫn chưa thực hiện. Tôi xin góp ý về phạm vi và thiết kế thực nghiệm trước giai đoạn cuối."),
    ]
    prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    for n,(title,bullets,notes) in enumerate(slides,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string("F6F8F3")
        for x,y,w,h,content,size,color in [(0.65,.45,12,1.4,title,30,"153D35"),(.85,2.2,11.7,4.1,"\n".join("• "+b for b in bullets),24,"183C35"),(.65,7.03,12,.3,f"Nguyễn Văn Duy · 2311553949     |     {n:02d}/12",11,"52665D")]:
            box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));box.text_frame.word_wrap=True
            for i,line in enumerate(content.split("\n")):
                p=box.text_frame.paragraphs[0] if i==0 else box.text_frame.add_paragraph();p.text=line;p.font.name="Arial";p.font.size=SlidePt(size);p.font.color.rgb=RGBColor.from_string(color);p.space_after=SlidePt(20)
        slide.notes_slide.notes_text_frame.text=notes
    prs.save(ROOT / "docs/slide-giua-ky-10-11-2026.pptx")
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        with tempfile.TemporaryDirectory(prefix="midterm-lo-") as tmp:
            subprocess.run([office,f"-env:UserInstallation=file://{tmp}","--headless","--convert-to","pdf","--outdir",str(ROOT / "docs"),str(stem.with_suffix('.docx')),str(ROOT / "docs/slide-giua-ky-10-11-2026.pptx")], check=True, stdout=subprocess.PIPE)
    print("Đã tạo báo cáo tiến độ và 12 slide giữa kỳ, kèm ghi chú trình bày.")


if __name__ == "__main__":
    main()
